import re
import json
import sys
from pathlib import Path
from typing import List, Dict, Any, Optional

# Ensure root directory is in python path so `database.*` / `src.*` imports resolve
# when this module is executed directly as a script.
ROOT_DIR = Path(__file__).resolve().parents[2]
if str(ROOT_DIR) not in sys.path:
    sys.path.append(str(ROOT_DIR))

from database.ingestor import ingest_to_postgres
from src.config import DOC_CLASS_PRS, DOC_CLASS_PTD, DOC_CLASS_RMM, DOC_CLASS_TEST
from src.engine1_ingestion.schema import IngestionOutput, ParsedItem, VisualEvidence


# --- Helper Functions ---

# Matches dotted engineering IDs such as "PR.UniGuide.FR.StartStop" or
# "PR.SmartNavigator.FR.DataIdentification.Patient.1", while ignoring version
# strings like "R3.1.5" (numeric segments are only allowed as a trailing index).
REQ_ID_PATTERN = re.compile(r"\b[A-Z][A-Za-z0-9]*(?:\.[A-Za-z][A-Za-z0-9]*){2,}(?:\.\d+)?\b")


def extract_doc_id(filename_or_path: str) -> str:
    """Extracts unique document ID like D001331681 using Regex, or falls back to stem name."""
    base_name = Path(filename_or_path).name
    match = re.search(r'(D\d{6,10})', base_name, re.IGNORECASE)
    if match:
        return match.group(1).upper()
    return Path(filename_or_path).stem


def extract_requirement_ids(text: str) -> List[str]:
    """Returns unique structured requirement IDs referenced in a block of text."""
    seen: Dict[str, None] = {}
    for match in REQ_ID_PATTERN.finditer(text or ""):
        seen.setdefault(match.group(0), None)
    return list(seen)


def classify_doc_class(path: Path, source_type: str) -> str:
    """Maps a source file/folder onto the logical doc class the linker queries."""
    if source_type == "TEST_FOLDER":
        return DOC_CLASS_TEST

    name = path.name.lower()
    if any(k in name for k in ("fmea", "risk", "hazard", "rmm")):
        return DOC_CLASS_RMM
    if any(k in name for k in ("test", "verification", "protocol")):
        return DOC_CLASS_TEST
    if any(k in name for k in ("design", "architecture", "ptd")):
        return DOC_CLASS_PTD
    return DOC_CLASS_PRS


# --- Parsers ---

# Header names that identify the primary key column of a tabular document
ID_HEADER_PATTERN = re.compile(
    r"\b(id|ids|identifier|req(uirement)?\s*(id|no|number)?|test\s*case|tc\s*id|item\s*(id|no))\b",
    re.IGNORECASE,
)


class DocumentParser:

    @staticmethod
    def parse_excel(file_path: str) -> IngestionOutput:
        import openpyxl

        path = Path(file_path)
        doc_id = extract_doc_id(file_path)
        doc_class = classify_doc_class(path, "EXCEL")
        item_type = "RISK_ITEM" if doc_class == DOC_CLASS_RMM else "REQUIREMENT"
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        items: List[ParsedItem] = []

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            # Find first non-empty row as header
            header_idx = -1
            headers = []
            for idx, row in enumerate(rows):
                non_empty = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if len(non_empty) >= 2:  # Assume table header has at least 2 columns
                    header_idx = idx
                    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(row)]
                    break

            if header_idx == -1:
                continue

            # Process data rows below header
            for row_num, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
                if not any(row):
                    continue  # Skip empty rows

                row_data = {}
                for col_idx, cell_value in enumerate(row):
                    if col_idx < len(headers) and cell_value is not None:
                        row_data[headers[col_idx]] = str(cell_value).strip()

                if not row_data:
                    continue

                # Identify Item ID or fallback to Row Number
                item_id = None
                for key, val in row_data.items():
                    if val and ID_HEADER_PATTERN.search(key):
                        item_id = val
                        break
                if not item_id:
                    # Fall back to any cell that is itself a structured requirement ID
                    item_id = next(
                        (v for v in row_data.values() if v and REQ_ID_PATTERN.fullmatch(v)),
                        None,
                    )
                if not item_id:
                    item_id = f"{doc_id}_{sheet_name}_R{row_num}"

                # Text payload aggregation
                text_payload = " | ".join([f"{k}: {v}" for k, v in row_data.items()])

                referenced_ids = [rid for rid in extract_requirement_ids(text_payload) if rid != item_id]

                items.append(ParsedItem(
                    item_id=str(item_id),
                    parent_doc_id=doc_id,
                    section_path=f"{doc_id} / {sheet_name}",
                    item_type=item_type,
                    text_content=text_payload,
                    metadata={**row_data, "referenced_ids": referenced_ids}
                ))

        wb.close()

        return IngestionOutput(
            document_id=doc_id,
            source_filename=path.name,
            source_type="EXCEL",
            doc_class=doc_class,
            total_items_extracted=len(items),
            items=items
        )

    @staticmethod
    def parse_docling(file_path: str) -> IngestionOutput:
        try:
            from docling.document_converter import DocumentConverter
        except (ImportError, OSError) as exc:
            raise RuntimeError(
                "DOCX/PDF parsing requires a working Docling/PyTorch installation. "
                "The current environment could not load Docling's native dependencies. "
                "Repair or reinstall the compatible PyTorch package, then retry."
            ) from exc

        doc_id = extract_doc_id(file_path)
        converter = DocumentConverter()
        result = converter.convert(file_path)
        doc = result.document

        items: List[ParsedItem] = []
        current_section = doc_id

        # Export converted markdown structure
        md_text = doc.export_to_markdown()
        lines = md_text.split('\n')

        chunk_buf: List[str] = []
        chunk_title = "General"

        def flush_chunk(title: str, section: str, body_lines: List[str]) -> None:
            """Emits a chunk, promoting it to a REQUIREMENT when the heading carries a real ID."""
            body = "\n".join(body_lines).strip()
            if not body:
                return

            heading_ids = extract_requirement_ids(title)
            if heading_ids:
                item_id = heading_ids[0]
                item_type = "REQUIREMENT"
            else:
                item_id = f"{doc_id}_SEC_{len(items) + 1:03d}"
                item_type = "SECTION"

            referenced_ids = [rid for rid in extract_requirement_ids(body) if rid != item_id]

            items.append(ParsedItem(
                item_id=item_id,
                parent_doc_id=doc_id,
                section_path=f"{doc_id} / {section}",
                item_type=item_type,
                title=title,
                text_content=body,
                metadata={"referenced_ids": referenced_ids}
            ))

        for line in lines:
            line_str = line.strip()
            if line_str.startswith("#"):
                flush_chunk(chunk_title, current_section, chunk_buf)
                chunk_buf = []

                # Update heading context
                current_section = line_str.lstrip("#").strip()
                chunk_title = current_section
            else:
                if line_str:
                    chunk_buf.append(line_str)

        # Flush final chunk
        flush_chunk(chunk_title, current_section, chunk_buf)

        path = Path(file_path)
        ext = path.suffix.replace(".", "").upper()
        return IngestionOutput(
            document_id=doc_id,
            source_filename=path.name,
            source_type=ext,
            doc_class=classify_doc_class(path, ext),
            total_items_extracted=len(items),
            items=items
        )

    @staticmethod
    def parse_test_folder(folder_path: str) -> IngestionOutput:
        folder = Path(folder_path)
        doc_id = extract_doc_id(folder.name)
        items: List[ParsedItem] = []

        # Scan for subdirectories (test cases) or treat root as one test case
        subfolders = [f for f in folder.iterdir() if f.is_dir()]
        target_folders = subfolders if subfolders else [folder]

        for tc_dir in target_folders:
            tc_id = tc_dir.name
            evidence_list: List[VisualEvidence] = []
            missing_artifacts: List[str] = []
            metadata: Dict[str, Any] = {}
            result_payload = ""

            # 1. Parse JSON result if present
            json_files = list(tc_dir.glob("*.json"))
            if json_files:
                try:
                    with open(json_files[0], 'r', encoding='utf-8') as jf:
                        metadata = json.load(jf)
                        result_payload = json.dumps(metadata, indent=2)
                except Exception as e:
                    metadata["json_error"] = str(e)
            else:
                missing_artifacts.append("result_json")

            # Keep the test case name in the embedded text so semantic search can use it
            text_content = f"Test Case: {tc_id}\n{result_payload}".strip()

            # 2. Fuzzy match screenshots/images
            image_extensions = ["*.png", "*.jpg", "*.jpeg"]
            image_files = []
            for ext in image_extensions:
                image_files.extend(tc_dir.glob(ext))

            found_actual = False
            found_diff = False

            for img in image_files:
                img_name = img.name.lower()
                rel_path = str(img)

                if any(k in img_name for k in ["diff", "pixel", "delta"]):
                    evidence_list.append(VisualEvidence(artifact_type="pixel_diff", file_path=rel_path))
                    found_diff = True
                elif any(k in img_name for k in ["annotated", "mark", "draw"]):
                    evidence_list.append(VisualEvidence(artifact_type="annotated_ss", file_path=rel_path))
                else:
                    evidence_list.append(VisualEvidence(artifact_type="actual_ss", file_path=rel_path))
                    found_actual = True

            if not found_actual:
                missing_artifacts.append("actual_screenshot")
            if not found_diff:
                missing_artifacts.append("pixel_diff_image")

            referenced_ids = extract_requirement_ids(f"{tc_id} {text_content}")

            items.append(ParsedItem(
                item_id=tc_id,
                parent_doc_id=doc_id,
                section_path=f"{doc_id} / {tc_id}",
                item_type="TEST_CASE",
                text_content=text_content,
                metadata={**metadata, "referenced_ids": referenced_ids},
                visual_evidence=evidence_list,
                missing_artifacts=missing_artifacts
            ))

        return IngestionOutput(
            document_id=doc_id,
            source_filename=folder.name,
            source_type="TEST_FOLDER",
            doc_class=DOC_CLASS_TEST,
            total_items_extracted=len(items),
            items=items
        )


# --- Universal Pipeline Runner ---

def parse_and_save(target_path: str, doc_class: Optional[str] = None) -> str:
    """Detects path type, parses, writes JSON, and pushes directly to PostgreSQL."""
    path = Path(target_path)
    if not path.exists():
        raise FileNotFoundError(f"Provided path does not exist: {target_path}")

    # Select parser based on input type
    if path.is_dir():
        result = DocumentParser.parse_test_folder(str(path))
        output_json_path = path / f"{result.document_id}_parsed.json"
    else:
        ext = path.suffix.lower()
        if ext in [".xlsx", ".xls"]:
            result = DocumentParser.parse_excel(str(path))
        elif ext in [".docx", ".pdf"]:
            result = DocumentParser.parse_docling(str(path))
        else:
            raise ValueError(f"Unsupported file extension: {ext}")

        output_json_path = path.parent / f"{path.stem}_parsed.json"

    # Allow the caller to override the auto-detected classification
    if doc_class:
        result.doc_class = doc_class

    # 1. Dump Pydantic payload to Dict
    parsed_dict_data = result.model_dump()

    # 2. Export local JSON for dev/debugging
    with open(output_json_path, "w", encoding="utf-8") as f:
        json.dump(parsed_dict_data, f, indent=2)

    print(f"✅ Successfully parsed '{path.name}' as {result.doc_class}")
    print(f"📊 Extracted Items: {result.total_items_extracted}")
    print(f"💾 Saved Output JSON to: {output_json_path}")

    # 3. AUTO-INGEST DIRECTLY TO POSTGRESQL
    ingest_to_postgres(str(output_json_path), parsed_dict_data)

    return str(output_json_path)


# --- Execution Hook ---

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python -m src.engine1_ingestion.parser <file-or-folder> [more paths...]")
        sys.exit(1)

    for target in sys.argv[1:]:
        target_path = Path(target)
        if not target_path.exists():
            print(f"⚠️ Skipping missing path: {target}")
            continue
        parse_and_save(str(target_path))